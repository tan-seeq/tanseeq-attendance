#!/usr/bin/env python3
"""
Replace October 2025 attendance data from an Excel file (Sheet1 with Arabic headers).
- Header expectations (Arabic): الموظف، التاريخ، الحضور، الانصراف، ساعات العمل، الحالة
- Works with English synonyms too: Employee, Date, Check-in, Check-out, Work Hours, Status
- Replaces data for cycle window: 2025-09-29 → 2025-10-28 (inclusive)

USAGE (inside container):
  python /app/backend/import_attendance_excel.py --url <EXCEL_URL>

Notes:
- Uses backend/.env (MONGO_URL, DB_NAME)
- Name mapping for known differences between Excel names and DB names is built-in
"""
import argparse
import asyncio
import io
import os
import sys
import uuid
from datetime import date, datetime, time
from pathlib import Path

import aiohttp
from dotenv import load_dotenv
from motor.motor_asyncio import AsyncIOMotorClient
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

# Ensure backend env
ROOT = Path(__file__).parent
load_dotenv(ROOT / ".env")

MONGO_URL = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.environ.get("DB_NAME", "tanseeq_hr")

# Cycle window for October 2025
CYCLE_START = date(2025, 9, 29)
CYCLE_END = date(2025, 10, 28)

# Name mapping from Excel → DB
NAME_MAPPING = {
    # Canonical English → DB names used in system
    "TAREK ABDELMONEM ZAKI ALWAZAN": "Tarek Wazzan",
    "TAREK ABDELMONEM ZAKI ALWAZAN ": "Tarek Wazzan",
    "KARIM MOHAMED MOUSTAFA ABDELMEGEUID": "Kareem",
    "GEHAD MOHAMED KAMAL AHMED": "Jihad",
    "HESHAM AHMED MOHAMED MOSTAFA": "Hesham",
    "MOHAMED AHMED MOHAMED MOSTAFA": "Mohamed Mostafa",
    "MOHAMED AHMED MOSTAFA": "Mohamed Mostafa",
    "Mohamed AHMED MOHAMED MOSTAFA": "Mohamed Mostafa",
    "Hatem Mohamed Ahmed": "Hatem Mohamed Ahmed",
}

# Column headers (Arabic/English)
HEADERS_MAP = {
    "employee": {"الموظف", "employee", "اسم الموظف"},
    "date": {"التاريخ", "date"},
    "check_in": {"الحضور", "check in", "check-in", "time in", "in"},
    "check_out": {"الانصراف", "check out", "check-out", "time out", "out"},
    "status": {"الحالة", "status"},
    "work_hours": {"ساعات العمل", "work hours"},
}


def normalize_str(s: str) -> str:
    return (s or "").strip().lower()


def excel_to_py_date(value, wb):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    # Excel serial number
    if isinstance(value, (int, float)):
        try:
            return from_excel(value, wb.epoch).date()
        except Exception:
            return None
    # String parse (YYYY-MM-DD or DD/MM/YYYY)
    try:
        return datetime.fromisoformat(str(value)).date()
    except Exception:
        pass
    try:
        return datetime.strptime(str(value), "%d/%m/%Y").date()
    except Exception:
        return None


def excel_time_to_str(value, wb) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.time().strftime("%H:%M:%S")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, (int, float)):
        try:
            dt = from_excel(value, wb.epoch)
            return dt.time().strftime("%H:%M:%S")
        except Exception:
            pass
    # Possibly HH:MM or HH:MM:SS string
    s = str(value).strip()
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(s, fmt).time().strftime("%H:%M:%S")
        except Exception:
            continue
    return None


async def download_excel(url: str) -> bytes:
    async with aiohttp.ClientSession() as session:
        async with session.get(url) as resp:
            resp.raise_for_status()
            return await resp.read()


async def main(url: str):
    print("📥 Downloading Excel from:", url)
    data = await download_excel(url)

    # Load workbook
    wb = load_workbook(filename=io.BytesIO(data), data_only=True)
    sheet = wb[wb.sheetnames[0]]
    print("📄 Using sheet:", sheet.title)

    # Try to locate header row within first 5 rows if row 1 fails
    def find_header_row(max_scan: int = 5):
        for r in range(1, max_scan + 1):
            row_cells = [str(c.value or "").strip() for c in next(sheet.iter_rows(min_row=r, max_row=r))]
            if not any(row_cells):
                continue
            # heuristic: contains at least الموظف and التاريخ
            if ("الموظف" in row_cells) and ("التاريخ" in row_cells):
                return r, row_cells
        # default to row 1
        return 1, [str(c.value or "").strip() for c in next(sheet.iter_rows(min_row=1, max_row=1))]

    header_row_idx, header_cells = find_header_row(max_scan=5)
    print(f"🧭 Detected header row at index {header_row_idx}:", header_cells)

    # Reset iterator after scanning
    # (openpyxl generators are stateless for iter_rows; we will use values_only later)

    # Build column index map (case-insensitive, Arabic-aware)
    col_idx = {}
    for i, hdr in enumerate(header_cells):
        hnorm = normalize_str(hdr)
        for key, aliases in HEADERS_MAP.items():
            if hnorm in aliases and key not in col_idx:
                col_idx[key] = i
    # Fallback matching for Arabic headers if normalization missed
    if "employee" not in col_idx:
        for i, hdr in enumerate(header_cells):
            if hdr.strip() == "الموظف":
                col_idx["employee"] = i
                break
    if "date" not in col_idx:
        for i, hdr in enumerate(header_cells):
            if hdr.strip() == "التاريخ":
                col_idx["date"] = i
                break
    # Not strictly required: check_in/check_out may be absent in user file
    if "check_in" not in col_idx:
        for i, hdr in enumerate(header_cells):
            if hdr.strip() == "الحضور":
                col_idx["check_in"] = i
                break
    if "check_out" not in col_idx:
        for i, hdr in enumerate(header_cells):
            if hdr.strip() == "الانصراف":
                col_idx["check_out"] = i
                break
    required = ["employee", "date"]
    for r in required:
        if r not in col_idx:
            raise RuntimeError(f"Missing required column in Excel header: {r}")

    # Connect DB
    client = AsyncIOMotorClient(MONGO_URL)
    db = client[DB_NAME]

    # Build employee name → id map (lowercased)
    users = await db.users.find({"is_active": True}).to_list(None)
    name_to_id = {normalize_str(u.get("name", "")): u["id"] for u in users}

    # Stage: delete existing attendance for cycle window
    start_str, end_str = CYCLE_START.isoformat(), CYCLE_END.isoformat()
    print(f"🧹 Deleting existing attendance between {start_str} and {end_str} (inclusive)…")
    delete_res = await db.attendance.delete_many({
        "date": {"$gte": start_str, "$lte": end_str}
    })
    print(f"   → Deleted {delete_res.deleted_count} records")

    # Insert new rows
    inserted = 0
    unknown_employees: set[str] = set()

    for row in sheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
        emp_raw = (row[col_idx["employee"]] if col_idx.get("employee") is not None else "")
        dt_raw = (row[col_idx["date"]] if col_idx.get("date") is not None else None)
        if not emp_raw:
            continue

        # Normalize employee name and map to DB
        excel_emp = str(emp_raw).strip()
        db_lookup = NAME_MAPPING.get(excel_emp.upper(), excel_emp)
        emp_id = name_to_id.get(normalize_str(db_lookup))
        if not emp_id:
            unknown_employees.add(excel_emp)
            continue

        # Parse date and filter by cycle
        d = excel_to_py_date(dt_raw, wb)
        if not d:
            continue
        if not (CYCLE_START <= d <= CYCLE_END):
            continue

        # Parse check-in/out
        ci = excel_time_to_str(row[col_idx["check_in"]] if col_idx.get("check_in") is not None else None, wb)
        co = excel_time_to_str(row[col_idx["check_out"]] if col_idx.get("check_out") is not None else None, wb)

        # Status
        status_val = str(row[col_idx["status"]]).strip() if col_idx.get("status") is not None and row[col_idx["status"]] is not None else ""
        status = "present"
        if not ci and not co:
            status = "absent" if status_val == "" or "غيا" in status_val or status_val.lower() == "absent" else "present"
        elif status_val:
            if "غيا" in status_val or status_val.lower() == "absent":
                status = "absent"
            elif "متأخ" in status_val or status_val.lower() == "late":
                status = "late"

        # Compute late minutes and early departure for storage
        late_minutes = 0
        early_departure_minutes = 0
        if ci:
            try:
                ci_t = datetime.strptime(ci, "%H:%M:%S").time()
                if ci_t > time(9, 0):
                    late_minutes = int((datetime.combine(d, ci_t) - datetime.combine(d, time(9, 0))).total_seconds() / 60)
            except Exception:
                pass
        if co:
            try:
                co_t = datetime.strptime(co, "%H:%M:%S").time()
                if co_t < time(18, 0):
                    early_departure_minutes = int((datetime.combine(d, time(18, 0)) - datetime.combine(d, co_t)).total_seconds() / 60)
            except Exception:
                pass

        record = {
            "id": str(uuid.uuid4()),
            "user_id": emp_id,
            "user_name": db_lookup,
            "date": d.isoformat(),
            "check_in": ci,
            "check_out": co,
            "status": status,
            "is_late": bool(late_minutes > 0),
            "late_minutes": late_minutes,
            "early_departure_minutes": early_departure_minutes,
            "deducted_hours": 0,
            "working_hours": None,
            "created_at": datetime.now().isoformat(),
        }

        await db.attendance.insert_one(record)
        inserted += 1

    print(f"✅ Inserted {inserted} attendance records in cycle window")
    if unknown_employees:
        print("⚠️ Unmapped employee names from Excel (please add to NAME_MAPPING if needed):")
        for n in sorted(unknown_employees):
            print("   -", n)

    print("🎯 Import complete.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--url", required=True, help="Public URL to Excel file")
    args = parser.parse_args()
    asyncio.run(main(args.url))
