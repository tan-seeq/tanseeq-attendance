#!/usr/bin/env python3
"""
Production Importer (Calibration October 2025)
- Upsert-only importer for TANSEEQ_attendance_report_2025-10-01_2025-10-27.xlsx
- Headers (English): Employee name, Date, Check In, Check Out, Working hours/duration, Status/Notes
- Window: 2025-09-29 → 2025-10-28 (inclusive)
- No deletes. Only update/insert attendance per (user_id, date)
- Sets status to "absent" only when the Status/Notes explicitly indicates absence
"""
import argparse
import asyncio
import io
import os
import re
from datetime import datetime, date, time
from pathlib import Path

import aiohttp
from dotenv import load_dotenv
from motor.motor_asyncio import AsyncIOMotorClient
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

ROOT = Path(__file__).parent
load_dotenv(ROOT / ".env")

MONGO_URL = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.environ.get("DB_NAME", "tanseeq_hr")

CYCLE_START = date(2025, 9, 29)
CYCLE_END = date(2025, 10, 28)

NAME_MAPPING = {
    "TAREK ABDELMONEM ZAKI ALWAZAN": "Tarek Wazzan",
    "KARIM MOHAMED MOUSTAFA ABDELMEGEUID": "Kareem",
    "GEHAD MOHAMED KAMAL AHMED": "Jihad",
    "HESHAM AHMED MOHAMED MOSTAFA": "Hesham",
    "MOHAMED AHMED MOHAMED MOSTAFA": "Mohamed Mostafa",
    "Mohamed AHMED MOHAMED MOSTAFA": "Mohamed Mostafa",
    "Hatem Mohamed Ahmed": "Hatem Mohamed Ahmed",
}

HEADERS = [
    "employee name", "date", "check in", "check out", "working hours/duration", "status/notes"
]


def norm(s: str) -> str:
    if s is None:
        return ""
    # normalize NBSP and multiple spaces
    x = str(s).replace("\xa0", " ")
    x = " ".join(x.split())
    return x.strip().lower()


def parse_excel_time(value, wb):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.time().strftime("%H:%M:%S")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, (int, float)):
        try:
            dt = from_excel(value, wb.epoch)
            return dt.time().strftime("%H:%M:%S")
        except Exception:
            pass
    s = str(value).strip()
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(s, fmt).time().strftime("%H:%M:%S")
        except Exception:
            continue
    return None


def parse_excel_date(value, wb):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            return from_excel(value, wb.epoch).date()
        except Exception:
            return None
    # iso or dd/mm/yyyy
    try:
        return datetime.fromisoformat(str(value)).date()
    except Exception:
        pass
    try:
        return datetime.strptime(str(value), "%d/%m/%Y").date()
    except Exception:
        return None


def parse_hours(value) -> float | None:
    if value is None:
        return None
    s = str(value).strip()
    m = re.search(r"(\d+(?:[\.,]\d+)?)", s)
    if not m:
        return None
    try:
        return float(m.group(1).replace(",", "."))
    except Exception:
        return None


async def download_excel(url: str) -> bytes:
    async with aiohttp.ClientSession() as session:
        async with session.get(url) as resp:
            resp.raise_for_status()
            return await resp.read()


async def main(url: str):
    print("📥 Downloading report:", url)
    data = await download_excel(url)
    wb = load_workbook(filename=io.BytesIO(data), data_only=True)
    sheet = wb[wb.sheetnames[0]]

    # detect header row
    header_row_idx = 1
    header_map = {}
    for r in range(1, 8):
        row = [str(c.value or "").strip() for c in next(sheet.iter_rows(min_row=r, max_row=r))]
        lowers = [norm(x) for x in row]
        # fuzzy detection for 'employee name'
        emp_candidates = [i for i, v in enumerate(lowers) if v in ("employee name", "employee", "employee name,", "employee name ")]
        date_candidates = [i for i, v in enumerate(lowers) if v in ("date", "date,")]
        if emp_candidates and date_candidates:
            header_row_idx = r
            header_map = {norm(v): i for i, v in enumerate(row)}
            # normalize keys if missing exact ones
            if "employee name" not in header_map and emp_candidates:
                header_map["employee name"] = emp_candidates[0]
            if "date" not in header_map and date_candidates:
                header_map["date"] = date_candidates[0]
            # optional keys
            for key, aliases in {
                "check in": ("check in", "checkin", "in"),
                "check out": ("check out", "checkout", "out"),
                "working hours/duration": ("working hours/duration", "working hours", "duration", "working hours " ),
                "status/notes": ("status/notes", "status", "notes", "status / notes"),
            }.items():
                if key not in header_map:
                    for i, v in enumerate(lowers):
                        if v in aliases:
                            header_map[key] = i
                            break
            break
    # ensure required
    required = ["employee name", "date"]
    for req in required:
        if req not in header_map:
            raise RuntimeError(f"Missing header: {req} -> found keys: {list(header_map.keys())}")

    client = AsyncIOMotorClient(MONGO_URL)
    db = client[DB_NAME]

    # build users map
    users = await db.users.find({"is_active": True}).to_list(None)
    name_to_id = {norm(u.get("name", "")): u["id"] for u in users}

    upserts = 0
    skipped = 0
    unknown = set()

    for row in sheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
        emp_raw = row[header_map["employee name"]]
        dt_raw = row[header_map["date"]]
        if not emp_raw:
            continue
        excel_name = str(emp_raw).strip()
        mapped_name = NAME_MAPPING.get(excel_name.upper(), excel_name)
        emp_id = name_to_id.get(norm(mapped_name))
        if not emp_id:
            unknown.add(excel_name)
            continue
        d = parse_excel_date(dt_raw, wb)
        if not d or not (CYCLE_START <= d <= CYCLE_END):
            continue

        ci = parse_excel_time(row[header_map.get("check in", -1)] if header_map.get("check in") is not None else None, wb)
        co = parse_excel_time(row[header_map.get("check out", -1)] if header_map.get("check out") is not None else None, wb)
        wh = parse_hours(row[header_map.get("working hours/duration", -1)] if header_map.get("working hours/duration") is not None else None)
        notes = str(row[header_map.get("status/notes", -1)] or "").strip()

        status = "present"
        is_absent = False
        if notes:
            if "absent" in notes.lower() or "غياب" in notes:
                status = "absent"
                is_absent = True
        # if both times missing but hours given, keep present with hours-only
        if not ci and not co and not is_absent and wh is not None:
            status = "present"

        late_minutes = 0
        if ci:
            try:
                ci_t = datetime.strptime(ci, "%H:%M:%S").time()
                if ci_t > time(9, 0):
                    late_minutes = int((datetime.combine(d, ci_t) - datetime.combine(d, time(9, 0))).total_seconds() / 60)
            except Exception:
                pass

        record = {
            "user_id": emp_id,
            "user_name": mapped_name,
            "date": d.isoformat(),
            "check_in": ci,
            "check_out": co,
            "working_hours": wh,
            "status": "absent" if is_absent else status,
            "is_late": bool(late_minutes > 0),
            "late_minutes": late_minutes,
            "early_departure_minutes": 0,
            "updated_at": datetime.now().isoformat(),
        }

        # Upsert by (user_id, date)
        res = await db.attendance.update_one(
            {"user_id": emp_id, "date": d.isoformat()},
            {"$set": record, "$setOnInsert": {"id": f"att-{emp_id}-{d.isoformat()}", "created_at": datetime.now().isoformat()}},
            upsert=True,
        )
        upserts += 1

    print(f"✅ Upserts: {upserts}")
    if unknown:
        print("⚠️ Unknown employees in report (please add to NAME_MAPPING if needed):")
        for n in sorted(unknown):
            print("   -", n)

    client.close()
    print("🎯 Import completed.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--url", required=True)
    args = parser.parse_args()
    asyncio.run(main(args.url))
